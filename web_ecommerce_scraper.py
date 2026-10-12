import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# 1. Requisição ao site de e-commerce
url = "http://books.toscrape.com/"
headers = {"User-Agent": "Mozilla/5.0"}
resposta = requests.get(url, headers=headers)
soup = BeautifulSoup(resposta.text, "html.parser")

# 2. Estruturação da planilha Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ofertas Livros"

# Cabeçalhos
headers_excel = ["Título do Livro", "Preço (£)", "Avaliação", "Status de Estoque"]
ws.append(headers_excel)

# Estilização do cabeçalho
cor_cabecalho = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fonte_cabecalho = Font(name="Arial", size=11, bold=True, color="FFFFFF")

for col in range(1, len(headers_excel) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = cor_cabecalho
    cell.font = fonte_cabecalho
    cell.alignment = Alignment(horizontal="center", vertical="center")

# 3. Extração dos produtos
livros = soup.find_all("article", class_="product_pod")

for livro in livros:
    titulo = livro.h3.a["title"]
    preco = livro.find("p", class_="price_color").text
    estoque = livro.find("p", class_="instock availability").text.strip()

    # Extração da classe da avaliação (ex: "star-rating Three")
    classes_avaliacao = livro.find("p", class_="star-rating")["class"]
    avaliacao = classes_avaliacao[1] if len(classes_avaliacao) > 1 else "N/A"

    ws.append([titulo, preco, avaliacao, estoque])

# Ajuste automático da largura das colunas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

# 4. Salvando o arquivo
nome_arquivo = "Relatorio_Ecommerce_Livros.xlsx"
wb.save(nome_arquivo)
print(f"Scraping concluído! Dados salvos em '{nome_arquivo}'.")